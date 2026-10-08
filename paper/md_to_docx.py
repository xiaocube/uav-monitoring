#!/usr/bin/env python3
"""将论文Markdown转换为Word文档（含封面、图片）"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
MD_PATH = ROOT / "paper" / "paper.md"
CHARTS_DIR = ROOT / "paper" / "charts"
OUTPUT_PATH = ROOT / "paper" / "论文_基于机器视觉的空域无人机目标识别监控方法研究.docx"


def set_chinese_font(run, font_name="宋体", size=10.5, bold=False):
    """设置中文字体"""
    font = run.font
    font.name = font_name
    font.size = Pt(size)
    font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)


def add_cover(doc):
    """添加论文封面"""
    for _ in range(3):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("项目研究报告")
    set_chinese_font(run, "黑体", 22, bold=True)

    doc.add_paragraph()
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("基于机器视觉的空域无人机目标识别监控方法研究")
    set_chinese_font(run, "黑体", 18, bold=True)

    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()

    info_items = [
        ("汇报人：", "________________________"),
        ("研究方向：", "机器视觉 / 低空安防"),
        ("项目周期：", "2026年6月 — 2026年7月"),
        ("汇报日期：", "2026年7月"),
    ]

    for label, value in info_items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_label = p.add_run(label)
        set_chinese_font(run_label, "宋体", 14)
        run_value = p.add_run(value)
        set_chinese_font(run_value, "宋体", 14)

    doc.add_page_break()


def parse_markdown(md_text):
    """解析Markdown文本为结构化数据"""
    lines = md_text.split('\n')
    result = []
    i = 0
    table_lines = []
    in_table = False

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 表格检测
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(stripped)
            i += 1
            continue
        elif in_table:
            # 表格结束
            if table_lines:
                result.append(('table', table_lines))
                table_lines = []
            in_table = False

        # 图片
        img_match = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', stripped)
        if img_match:
            alt_text = img_match.group(1)
            img_path = img_match.group(2)
            result.append(('image', alt_text, img_path))
            i += 1
            continue

        # 标题
        h_match = re.match(r'^(#{1,6})\s+(.+)$', stripped)
        if h_match:
            level = len(h_match.group(1))
            text = h_match.group(2)
            result.append(('heading', level, text))
            i += 1
            continue

        # 分隔线
        if stripped == '---':
            result.append(('separator',))
            i += 1
            continue

        # 普通段落（跳过空行和标记行）
        if stripped and not stripped.startswith('【配图提示'):
            result.append(('paragraph', stripped))

        i += 1

    if table_lines:
        result.append(('table', table_lines))

    return result


def add_table(doc, table_lines):
    """添加表格"""
    if not table_lines:
        return

    # 解析表头
    headers = [c.strip() for c in table_lines[0].split('|')[1:-1]]

    # 检查是否有分隔行（如 |---|---|）
    data_start = 1
    if len(table_lines) > 1 and re.match(r'^\|[\s\-:|]+\|$', table_lines[1]):
        data_start = 2

    # 解析数据行
    rows_data = []
    for line in table_lines[data_start:]:
        cells = [c.strip() for c in line.split('|')[1:-1]]
        if cells and any(cells):
            rows_data.append(cells)

    if not rows_data:
        return

    table = doc.add_table(rows=1 + len(rows_data), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 表头
    for j, header in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = header
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_chinese_font(run, "黑体", 10, bold=True)

    # 数据行
    for i, row_data in enumerate(rows_data):
        for j, cell_text in enumerate(row_data):
            if j < len(headers):
                cell = table.rows[i + 1].cells[j]
                cell.text = cell_text
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        set_chinese_font(run, "宋体", 10)

    doc.add_paragraph()


def add_image(doc, alt_text, img_path):
    """添加图片"""
    # 解析相对路径
    if img_path.startswith('paper/'):
        full_path = ROOT / img_path
    else:
        full_path = ROOT / img_path

    if not full_path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"[图片未找到: {img_path}]")
        set_chinese_font(run, "宋体", 10)
        return

    try:
        from PIL import Image
        with Image.open(full_path) as img:
            width_px, height_px = img.size

        # 计算图片显示尺寸（最大宽度14cm）
        max_width_cm = 14
        aspect = height_px / width_px
        display_width_cm = min(max_width_cm, width_px / 96 * 2.54)
        display_height_cm = display_width_cm * aspect

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(full_path), width=Cm(display_width_cm))

        # 图题
        p_caption = doc.add_paragraph()
        p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_caption = p_caption.add_run(alt_text)
        set_chinese_font(run_caption, "宋体", 10.5)

        doc.add_paragraph()
    except Exception as e:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"[图片插入失败: {img_path} - {e}]")
        set_chinese_font(run, "宋体", 10)


def add_heading(doc, level, text):
    """添加标题"""
    if level == 1:
        p = doc.add_heading(text, level=1)
        for run in p.runs:
            set_chinese_font(run, "黑体", 16, bold=True)
    elif level == 2:
        p = doc.add_heading(text, level=2)
        for run in p.runs:
            set_chinese_font(run, "黑体", 14, bold=True)
    elif level == 3:
        p = doc.add_heading(text, level=3)
        for run in p.runs:
            set_chinese_font(run, "黑体", 12, bold=True)
    else:
        p = doc.add_paragraph()
        run = p.add_run(text)
        set_chinese_font(run, "黑体", 11, bold=True)

    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_paragraph_text(doc, text):
    """添加正文段落"""
    # 处理加粗文本 **text**
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.5

    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            set_chinese_font(run, "黑体", 10.5, bold=True)
        else:
            run = p.add_run(part)
            set_chinese_font(run, "宋体", 10.5)


def add_separator(doc):
    """添加分隔线"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("—" * 30)
    set_chinese_font(run, "宋体", 10)
    doc.add_paragraph()


def set_document_styles(doc):
    """设置文档默认样式"""
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    # 页面设置
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)


def main():
    print("读取Markdown文件...")
    if not MD_PATH.exists():
        print(f"错误: 找不到 {MD_PATH}")
        sys.exit(1)

    with open(MD_PATH, 'r', encoding='utf-8') as f:
        md_text = f.read()

    print("解析Markdown...")
    parsed = parse_markdown(md_text)

    print("创建Word文档...")
    doc = Document()
    set_document_styles(doc)

    # 添加封面
    print("添加封面...")
    add_cover(doc)

    # 添加正文
    print("添加正文内容...")
    for item in parsed:
        kind = item[0]
        if kind == 'heading':
            add_heading(doc, item[1], item[2])
        elif kind == 'paragraph':
            add_paragraph_text(doc, item[1])
        elif kind == 'table':
            add_table(doc, item[1])
        elif kind == 'image':
            add_image(doc, item[1], item[2])
        elif kind == 'separator':
            add_separator(doc)

    # 保存
    print(f"保存Word文档: {OUTPUT_PATH}")
    doc.save(OUTPUT_PATH)
    print("✅ Word文档生成完成!")
    print(f"文件位置: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
